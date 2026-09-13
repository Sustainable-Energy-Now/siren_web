# powermatchui/views/esoo_export_views.py
"""
Web download of all WEM ESOO Foundation models (WS1) as an Excel workbook,
mirroring the export_facilities_excel.py / export_gridlines_excel.py
management-command pattern but streamed directly as a browser download,
matching powerplotui.views.facility_scada_views' export_scada_excel view.

Sheets produced:
  EsooVintages          - one row per WEM ESOO edition/tier
  EsooMethodVersions     - AEMO forecasting-method version timeline
  EsooTaxonomyMappings   - native scenario label -> Low/Expected/High mapping
  EsooFigures            - per-figure demand/supply-adequacy values + provenance
  EsooSourceDocuments    - retrieved source documents (SourceDocument rows
                           where esoo_vintage is set)
  AnnualDemandActuals    - SCADA-derived actuals used for G2 bias comparison
"""
from django.contrib.auth.decorators import login_required
from django.http import HttpResponse
from django.utils import timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from siren_web.models import (
    AnnualDemandActual,
    EsooFigure,
    EsooMethodVersion,
    EsooTaxonomyMapping,
    EsooVintage,
    SourceDocument,
)

HEADER_FONT = Font(bold=True, color='FFFFFF')
HEADER_FILL = PatternFill(start_color='2E75B6', end_color='2E75B6', fill_type='solid')
HEADER_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)


def _naive(dt):
    """openpyxl rejects tz-aware datetimes outright; USE_TZ=True means every
    auto_now/auto_now_add field comes back aware, so every datetime written
    to a cell must be converted to local time and stripped of tzinfo first."""
    if dt is None:
        return None
    if timezone.is_aware(dt):
        return timezone.localtime(dt).replace(tzinfo=None)
    return dt


def _write_sheet(ws, headers, rows):
    ws.append(headers)
    for cell in ws[1]:
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = HEADER_ALIGN
    for row in rows:
        ws.append(row)
    for col_idx, header in enumerate(headers, 1):
        col_letter = get_column_letter(col_idx)
        max_len = len(str(header))
        for row in ws.iter_rows(min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value is not None:
                    max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = min(max_len + 2, 60)


def _build_esoo_workbook() -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)

    # ------------------------------------------------------------------
    # EsooVintages
    # ------------------------------------------------------------------
    ws = wb.create_sheet('EsooVintages')
    headers = [
        'ID', 'Year', 'Tier', 'Publication Date', 'Source URL', 'Checksum',
        'Local File Path', 'Method Version', 'Ingestion Status', 'Notes',
        'Created At', 'Updated At',
    ]
    rows = []
    for v in EsooVintage.objects.select_related('method_version').order_by('year'):
        rows.append([
            v.idesoovintage,
            v.year,
            v.get_tier_display(),
            v.publication_date,
            v.source_url,
            v.checksum,
            v.local_file_path,
            v.method_version.version_label if v.method_version else None,
            v.get_ingestion_status_display(),
            v.notes,
            _naive(v.created_at),
            _naive(v.updated_at),
        ])
    _write_sheet(ws, headers, rows)

    # ------------------------------------------------------------------
    # EsooMethodVersions
    # ------------------------------------------------------------------
    ws = wb.create_sheet('EsooMethodVersions')
    headers = ['ID', 'Version Label', 'Effective From Vintage', 'Is Breaking Change', 'Description']
    rows = []
    for m in EsooMethodVersion.objects.order_by('effective_from_vintage'):
        rows.append([
            m.idesoomethodversion,
            m.version_label,
            m.effective_from_vintage,
            m.is_breaking_change,
            m.description,
        ])
    _write_sheet(ws, headers, rows)

    # ------------------------------------------------------------------
    # EsooTaxonomyMappings
    # ------------------------------------------------------------------
    ws = wb.create_sheet('EsooTaxonomyMappings')
    headers = ['ID', 'Vintage Year', 'Native Label', 'Mapped Scenario', 'Notes']
    rows = []
    for t in EsooTaxonomyMapping.objects.select_related('vintage').order_by('vintage__year', 'native_label'):
        rows.append([
            t.idesootaxonomymapping,
            t.vintage.year,
            t.native_label,
            t.get_mapped_scenario_display(),
            t.notes,
        ])
    _write_sheet(ws, headers, rows)

    # ------------------------------------------------------------------
    # EsooFigures
    # ------------------------------------------------------------------
    ws = wb.create_sheet('EsooFigures')
    headers = [
        'ID', 'Vintage Year', 'Domain', 'Metric', 'Forecast Year',
        'Demand Growth Scenario', 'POE Level', 'Demand Basis', 'Value', 'Unit',
        'Source Document', 'Source Version', 'Table Ref', 'Page Ref', 'Cell Ref',
        'Extraction Date', 'Extraction Method', 'Reconciliation Adjustment',
        'Validation Status', 'Validation Notes', 'Created At',
    ]
    rows = []
    for f in EsooFigure.objects.select_related('vintage').order_by(
        'vintage__year', 'domain', 'metric', 'forecast_year', 'demand_growth_scenario', 'poe_level'
    ):
        rows.append([
            f.idesoofigure,
            f.vintage.year,
            f.get_domain_display(),
            f.get_metric_display(),
            f.forecast_year,
            f.get_demand_growth_scenario_display(),
            f.poe_level,
            f.get_demand_basis_display(),
            f.value,
            f.unit,
            f.source_document,
            f.source_version,
            f.table_ref,
            f.page_ref,
            f.cell_ref,
            f.extraction_date,
            f.get_extraction_method_display(),
            f.reconciliation_adjustment,
            f.get_validation_status_display(),
            f.validation_notes,
            _naive(f.created_at),
        ])
    _write_sheet(ws, headers, rows)

    # ------------------------------------------------------------------
    # EsooSourceDocuments (SourceDocument rows backing an ESOO vintage)
    # ------------------------------------------------------------------
    ws = wb.create_sheet('EsooSourceDocuments')
    headers = [
        'ID', 'Vintage Year', 'Doc Type', 'Source URL', 'Checksum',
        'Local File Path', 'Retrieved At',
    ]
    rows = []
    for d in SourceDocument.objects.filter(esoo_vintage__isnull=False).select_related(
        'esoo_vintage'
    ).order_by('esoo_vintage__year', 'doc_type'):
        rows.append([
            d.idsourcedocument,
            d.esoo_vintage.year,
            d.get_doc_type_display(),
            d.source_url,
            d.checksum,
            d.local_file_path,
            _naive(d.retrieved_at),
        ])
    _write_sheet(ws, headers, rows)

    # ------------------------------------------------------------------
    # AnnualDemandActuals (SCADA-derived actuals, ESOO bias comparison)
    # ------------------------------------------------------------------
    ws = wb.create_sheet('AnnualDemandActuals')
    headers = [
        'ID', 'Year', 'Demand Basis', 'Annual Energy (GWh)',
        'Peak Demand (MW)', 'Peak Datetime', 'Minimum Demand (MW)',
        'Minimum Datetime', 'Computed At',
    ]
    rows = []
    for a in AnnualDemandActual.objects.order_by('year', 'demand_basis'):
        rows.append([
            a.idannualdemandactual,
            a.year,
            a.get_demand_basis_display(),
            a.annual_energy_gwh,
            a.peak_demand_mw,
            _naive(a.peak_datetime),
            a.minimum_demand_mw,
            _naive(a.minimum_datetime),
            _naive(a.computed_at),
        ])
    _write_sheet(ws, headers, rows)

    return wb


@login_required
def export_esoo_excel(request):
    """Stream all ESOO Foundation models as a single .xlsx download."""
    wb = _build_esoo_workbook()

    response = HttpResponse(
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    )
    response['Content-Disposition'] = 'attachment; filename="esoo_export.xlsx"'
    wb.save(response)
    return response
